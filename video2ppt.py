import sys
import os
import cv2
from pptx import Presentation
from pptx.util import Inches
from PyQt5.QtWidgets import (QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel, QFileDialog, QSlider, QStyle, QSizePolicy)
from PyQt5.QtGui import QPixmap, QImage
from PyQt5.QtCore import Qt, QTimer

class VideoCapture(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("MP4 to PPT Converter")
        self.setGeometry(100, 100, 800, 600)
        
        self.video_path = ""
        self.cap = None
        self.playing = False
        self.frames = []

        self.initUI()

    def initUI(self):
        layout = QVBoxLayout()

        # Video Player
        self.video_label = QLabel(self)
        self.video_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.video_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(self.video_label)

        # Controls
        controls_layout = QHBoxLayout()

        self.open_button = QPushButton("Open MP4", self)
        self.open_button.clicked.connect(self.open_file)
        controls_layout.addWidget(self.open_button)

        self.play_button = QPushButton("Play", self)
        self.play_button.clicked.connect(self.play_video)
        controls_layout.addWidget(self.play_button)

        self.pause_button = QPushButton("Pause", self)
        self.pause_button.clicked.connect(self.pause_video)
        controls_layout.addWidget(self.pause_button)

        self.capture_button = QPushButton("Capture Frame", self)
        self.capture_button.clicked.connect(self.capture_frame)
        controls_layout.addWidget(self.capture_button)

        self.ppt_button = QPushButton("Create PPT", self)
        self.ppt_button.clicked.connect(self.create_ppt)
        controls_layout.addWidget(self.ppt_button)

        self.status_label = QLabel("Captured Frames: 0", self)
        controls_layout.addWidget(self.status_label)

        layout.addLayout(controls_layout)

        # Slider
        self.slider = QSlider(Qt.Horizontal, self)
        self.slider.setRange(0, 100)
        self.slider.valueChanged.connect(self.slider_moved)
        layout.addWidget(self.slider)

        self.setLayout(layout)

        # Timer for updating frames
        self.timer = QTimer()
        self.timer.timeout.connect(self.update_frame)

    def open_file(self):
        options = QFileDialog.Options()
        self.video_path, _ = QFileDialog.getOpenFileName(self, "Open MP4 File", "", "MP4 Files (*.mp4);;All Files (*)", options=options)
        if self.video_path:
            self.cap = cv2.VideoCapture(self.video_path)
            self.slider.setRange(0, int(self.cap.get(cv2.CAP_PROP_FRAME_COUNT)) - 1)
            self.show_frame()

    def play_video(self):
        if self.cap:
            self.playing = True
            self.timer.start(30)

    def pause_video(self):
        self.playing = False
        self.timer.stop()

    def capture_frame(self):
        if hasattr(self, 'current_frame'):
            self.frames.append(self.current_frame)
            self.update_status_label()

    def update_status_label(self):
        self.status_label.setText(f"Captured Frames: {len(self.frames)}")

    def update_frame(self):
        if self.cap and self.playing:
            ret, frame = self.cap.read()
            if ret:
                self.current_frame = frame
                self.display_frame(frame)
                self.slider.setValue(int(self.cap.get(cv2.CAP_PROP_POS_FRAMES)))

    def slider_moved(self):
        if self.cap and not self.playing:
            frame_number = self.slider.value()
            self.cap.set(cv2.CAP_PROP_POS_FRAMES, frame_number)
            self.show_frame()

    def show_frame(self):
        ret, frame = self.cap.read()
        if ret:
            self.current_frame = frame
            self.display_frame(frame)

    def display_frame(self, frame):
        frame = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
        qimg = QImage(frame, frame.shape[1], frame.shape[0], QImage.Format_RGB888)
        pixmap = QPixmap.fromImage(qimg)
        self.video_label.setPixmap(pixmap.scaled(self.video_label.size(), Qt.KeepAspectRatio))

    def create_ppt(self):
        if not self.frames:
            return
        
        # Get the directory and name of the video file
        video_dir = os.path.dirname(self.video_path)
        video_name = os.path.splitext(os.path.basename(self.video_path))[0]
        
        # Set the filename for the PPT
        ppt_filename = os.path.join(video_dir, f"{video_name}.pptx")
        
        prs = Presentation()
        
        # Set slide size to 16:9 (13.33 inches x 7.5 inches)
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        target_width = 960  # Width in pixels (13.33 inches at 72 DPI)
        target_height = 540  # Height in pixels (7.5 inches at 72 DPI)

        for i, frame in enumerate(self.frames):
            # Calculate the aspect ratio
            h, w, _ = frame.shape
            aspect_ratio = w / h

            # Determine the new size while maintaining the aspect ratio
            if aspect_ratio > (target_width / target_height):
                new_width = target_width
                new_height = int(target_width / aspect_ratio)
            else:
                new_height = target_height
                new_width = int(target_height * aspect_ratio)

            # Resize the frame
            frame_resized = cv2.resize(frame, (new_width, new_height))

            # Save the resized frame as an image file
            image_path = os.path.join(video_dir, f'frame_{i}.jpg')
            cv2.imwrite(image_path, frame_resized)

            # Add a new slide and insert the image
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            left = Inches((13.33 - new_width / 72) / 2)  # Center horizontally on the 13.33 inch slide
            top = Inches((7.5 - new_height / 72) / 2)  # Center vertically on the 7.5 inch slide
            slide.shapes.add_picture(image_path, left, top, width=Inches(new_width / 72), height=Inches(new_height / 72))

            # Remove the image file to clean up
            os.remove(image_path)

        prs.save(ppt_filename)
        print(f"PPT saved at: {ppt_filename}")

if __name__ == '__main__':
    app = QApplication(sys.argv)
    window = VideoCapture()
    window.show()
    sys.exit(app.exec_())

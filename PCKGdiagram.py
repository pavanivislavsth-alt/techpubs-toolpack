import sys
import os
from openpyxl import load_workbook
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QLabel, QLineEdit, QPushButton, QVBoxLayout, QWidget, QGridLayout
)
from PyQt5.QtCore import Qt

class PackageDiagramSearchApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.initUI()
        self.load_workbook()

    def initUI(self):
        self.setWindowTitle("Package Diagram Number Search")
        self.setGeometry(100, 100, 600, 400)

        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        self.layout = QGridLayout()
        self.layout.setContentsMargins(20, 20, 20, 20)
        self.layout.setSpacing(10)

        self.entry_label = QLabel("Enter Cypress Package Diagram Number:")
        self.layout.addWidget(self.entry_label, 0, 0, 1, 3, Qt.AlignCenter)

        self.entry_field = QLineEdit()
        self.entry_field.setFixedWidth(300)
        self.layout.addWidget(self.entry_field, 1, 0, 1, 3, Qt.AlignCenter)

        self.search_button = QPushButton("Search")
        self.search_button.setFixedSize(100, 40)
        self.search_button.setStyleSheet("""
            QPushButton {
                background-color: #0A8276;
                color: white;
                border-radius: 10px;
            }
            QPushButton:hover {
                background-color: #087563;
            }
            QPushButton:pressed {
                background-color: #06584C;
            }
        """)
        self.search_button.clicked.connect(self.search_cypress_number)
        self.layout.addWidget(self.search_button, 2, 1, Qt.AlignCenter)

        self.result_label = QLabel("")
        self.result_label.setWordWrap(True)
        self.result_label.setStyleSheet("font-family: Helvetica; font-size: 12pt;")
        self.layout.addWidget(self.result_label, 3, 0, 1, 3, Qt.AlignCenter)

        self.layout.setRowStretch(0, 1)
        self.layout.setRowStretch(1, 1)
        self.layout.setRowStretch(2, 1)
        self.layout.setRowStretch(3, 5)
        self.layout.setColumnStretch(0, 1)
        self.layout.setColumnStretch(1, 1)
        self.layout.setColumnStretch(2, 1)

        central_widget.setLayout(self.layout)

    def load_workbook(self):
        wb_path = os.path.join(os.path.dirname(__file__), 'PLAN2.xlsx')
        self.wb = load_workbook(wb_path)
        self.ws = self.wb.active

    def search_cypress_number(self):
        cypress_number = self.entry_field.text().strip().upper()

        results = []

        for row in self.ws.iter_rows(values_only=True):
            if all(cell is not None for cell in row):
                excel_cypress_number = str(row[1]).strip().upper()
                if excel_cypress_number == cypress_number:
                    infineon_number = row[0]
                    results.append(infineon_number)

        if results:
            result_text = "Infineon Package Diagram Numbers:\n" + "\n".join(results)
            self.result_label.setText(result_text)
            self.result_label.setStyleSheet("color: green; font-family: Helvetica; font-size: 12pt;")
        else:
            self.result_label.setText("Not found")
            self.result_label.setStyleSheet("color: red; font-family: Helvetica; font-size: 12pt;")

def main():
    app = QApplication(sys.argv)
    ex = PackageDiagramSearchApp()
    ex.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()
